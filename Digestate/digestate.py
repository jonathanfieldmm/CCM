import openpyxl

# Load the workbook and select the 'Sites' sheet
input_wb = openpyxl.load_workbook('Digestate/Anaerobic-Digestion-Deployment-in-the-United-Kingdom-April-2024-Operational-RAW.xlsx')
input_sheet = input_wb['Sites']

# Create a new workbook and select the active sheet
output_wb = openpyxl.Workbook()
output_sheet = output_wb.active

# Iterate through each row in the input sheet and write it to the output sheet
for row in input_sheet.iter_rows(values_only=True):
    output_sheet.append(row)

# Save the new workbook
output_wb.save('Digestate/Anaerobic-Digestion-Deployment-Copy.xlsx')

# Print a success message
print("Data from the 'Sites' sheet has been successfully copied to 'Anaerobic-Digestion-Deployment-Copy.xlsx'")
